"""Render a PPT deck from a slide spec built by lib/ppt-spec.ts.

All numbers, chart data, colours and explanation text are computed on the
TypeScript side with the same helpers the web dashboard uses; this script only
lays them out on the SIG template (Black/Light).

stdin: JSON spec {theme, slides: [...]}. argv[1]: output .pptx path.
"""
import sys, os, json
# python-pptx & deps are vendored in lib/python-packages. Add them to sys.path
# here: embedded interpreters (._pth file) ignore PYTHONPATH.
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "python-packages"))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATES = {
    "black": os.path.join(HERE, "ppt-templates", "sig-black.pptx"),
    "light": os.path.join(HERE, "ppt-templates", "sig-light.pptx"),
}
FONT = "SIG Text"


def rgb(hex_color):
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# Web palette (components/charts/colors.ts). Content slides are light in both
# SIG templates, so body styling is shared; only the title slide differs.
TEAL = rgb("#174D55")
TURQUOISE = rgb("#4ECDC4")
MINT = rgb("#F7FFF7")
WHITE = rgb("#FFFFFF")
TEXT = rgb("#0F172A")
MUTED = rgb("#4E8088")
BORDER = rgb("#C9E4E1")
GRID = rgb("#E3EEED")
BANNER_FILL = rgb("#FFF8D6")
BANNER_LINE = rgb("#FFE66D")

TITLE_TEXT = {
    "black": (rgb("#F5F7FA"), rgb("#CBD5E1")),
    "light": (rgb("#0F172A"), rgb("#475569")),
}

# Content area of the "Title Only" layout (inches): below the title placeholder,
# above the slide-number footer.
LEFT, RIGHT, TOP, BOTTOM = 0.5, 12.83, 1.45, 6.85
GAP = 0.15


# ----------------------------------------------------------------------------
# text / shape helpers
# ----------------------------------------------------------------------------
def style_run(run, size, bold=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def text_box(slide, text, x, y, w, h, size=10, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    style_run(p.add_run(), size, bold, color)
    p.runs[0].text = text
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def get_ph(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def delete_all_slides(prs):
    """Drop the template's built-in sample slides so output starts clean."""
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.get(qn("r:id")))
        lst.remove(sid)


def truncate(label, n=30):
    label = str(label)
    return label if len(label) <= n else label[: n - 1] + "…"


# ----------------------------------------------------------------------------
# building blocks
# ----------------------------------------------------------------------------
def kpi_strip(slide, kpis, y):
    """KPI cards like the web's ExportKpi/ReportKpi. Returns the y below them."""
    if not kpis:
        return y
    per_row = len(kpis) if len(kpis) <= 5 else 4
    card_h = 0.82 if len(kpis) > 5 else 0.92
    width = RIGHT - LEFT
    for i, kpi in enumerate(kpis):
        row, col = divmod(i, per_row)
        cols = min(per_row, len(kpis) - row * per_row)
        w = (width - GAP * (per_row - 1)) / per_row
        x = LEFT + col * (w + GAP)
        cy = y + row * (card_h + 0.1)
        rect(slide, x, cy, w, card_h, WHITE, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, x + 0.06, cy + 0.14, 0.05, card_h - 0.28, TURQUOISE)
        text_box(slide, kpi["label"].upper(), x + 0.18, cy + 0.06, w - 0.25, 0.25, 8, True, MUTED)
        value = kpi["value"]
        size = 16 if len(value) <= 14 else 13 if len(value) <= 22 else 11
        text_box(slide, value, x + 0.18, cy + 0.28, w - 0.25, 0.36, size, True, TEAL)
        if kpi.get("sub"):
            text_box(slide, kpi["sub"], x + 0.18, cy + card_h - 0.3, w - 0.25, 0.24, 8, False, MUTED)
    rows = (len(kpis) + per_row - 1) // per_row
    return y + rows * card_h + (rows - 1) * 0.1 + GAP


def banner(slide, text, y):
    rect(slide, LEFT, y, RIGHT - LEFT, 0.42, BANNER_FILL, BANNER_LINE)
    text_box(slide, text, LEFT + 0.12, y, RIGHT - LEFT - 0.24, 0.42, 9, False, TEAL, anchor=MSO_ANCHOR.MIDDLE)
    return y + 0.42 + GAP


def explain_box(slide, title, bullets, x, y, w, h, paragraph=None, size=None, spacing=6):
    """'Penjelasan' panel: header + bullet list (and an optional paragraph)."""
    rect(slide, x, y, w, h, MINT, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, x, y, w, 0.34, TEAL)
    text_box(slide, title, x + 0.12, y, w - 0.24, 0.34, 9, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

    length = len(paragraph or "") + sum(len(b) for b in bullets)
    area = w * (h - 0.4)
    if size is None:
        size = 11 if length < area * 40 else 10 if length < area * 55 else 9

    box = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.42), Inches(w - 0.2), Inches(h - 0.5))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    if paragraph:
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = paragraph
        style_run(r, size, False, TEXT)
        p.space_after = Pt(10)
        first = False
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = f"•  {b}"
        style_run(r, size, False, TEXT)
        p.space_after = Pt(spacing)


def panel(slide, title, x, y, w, h):
    """Chart panel like the web's ChartPanel: teal header bar + white body."""
    rect(slide, x, y, w, h, WHITE, BORDER)
    rect(slide, x, y, w, 0.32, TEAL)
    text_box(slide, title.upper(), x + 0.1, y, w - 0.2, 0.32, 8, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    return x + 0.08, y + 0.38, w - 0.16, h - 0.44


# ----------------------------------------------------------------------------
# charts
# ----------------------------------------------------------------------------
def _no_fill():
    return parse_xml('<c:spPr %s><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>' % nsdecls("c", "a"))


def base_style(chart, legend):
    cs = chart._chartSpace
    chart_el = cs.find(qn("c:chart"))
    chart_el.addnext(_no_fill())
    plot_area = chart_el.find(qn("c:plotArea"))
    if plot_area is not None:
        plot_area.append(_no_fill())
    chart.has_title = False
    chart.has_legend = legend
    chart.font.name = FONT
    chart.font.size = Pt(8)
    chart.font.color.rgb = TEXT
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(8)
        chart.legend.font.color.rgb = TEAL


def style_axes(chart, value_format=None, show_values=True, min_zero=False):
    cat = chart.category_axis
    cat.tick_labels.font.size = Pt(8)
    cat.tick_labels.font.color.rgb = TEXT
    cat.format.line.color.rgb = GRID
    cat.has_major_gridlines = False
    val = chart.value_axis
    val.visible = show_values
    val.has_major_gridlines = show_values
    if show_values:
        val.major_gridlines.format.line.color.rgb = GRID
        val.tick_labels.font.size = Pt(7)
        val.tick_labels.font.color.rgb = MUTED
        val.format.line.fill.background()
        if min_zero:
            val.minimum_scale = 0
        if value_format:
            val.tick_labels.number_format = value_format
            val.tick_labels.number_format_is_linked = False


def add_hbar(slide, spec, x, y, w, h):
    cd = CategoryChartData()
    cd.categories = [truncate(c) for c in spec["categories"]]
    cd.add_series("Nilai", spec["values"])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    base_style(chart, legend=False)
    style_axes(chart, show_values=False)
    chart.category_axis.reverse_order = True  # first (largest) item on top, like the web
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW  # stay left of negative bars
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.series[0].invert_if_negative = False
    color, negative = rgb(spec["color"]), rgb(spec["negativeColor"])
    for i, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = negative if spec["values"][i] < 0 else color
        dl = point.data_label
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        tf = dl.text_frame
        tf.text = spec["labels"][i]
        style_run(tf.paragraphs[0].runs[0], 8, True, TEAL)
    # Per-point override: without it PowerPoint draws negative bars hollow.
    for dpt in plot.series[0]._element.findall(qn("c:dPt")):
        if dpt.find(qn("c:invertIfNegative")) is None:
            dpt.find(qn("c:idx")).addnext(parse_xml('<c:invertIfNegative %s val="0"/>' % nsdecls("c")))


def add_line(slide, spec, x, y, w, h):
    cd = CategoryChartData()
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["name"], s["values"])
    many = len(spec["categories"]) > 18
    kind = XL_CHART_TYPE.LINE if many else XL_CHART_TYPE.LINE_MARKERS
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    base_style(chart, legend=len(spec["series"]) > 1)
    values = [v for s in spec["series"] for v in s["values"] if v is not None]
    style_axes(chart, spec.get("axisFormat"), min_zero=bool(values) and min(values) >= 0)
    plot = chart.plots[0]
    for s_spec, series in zip(spec["series"], plot.series):
        color = rgb(s_spec["color"])
        series.smooth = True
        series.format.line.color.rgb = color
        series.format.line.width = Pt(3 if s_spec.get("emphasis") else 2.25)
        if s_spec.get("dashed"):
            series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if not many:
            series.marker.size = 5
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.color.rgb = color
    last_label = spec.get("lastLabel")
    if last_label and plot.series and spec["categories"]:
        dl = plot.series[0].points[len(spec["categories"]) - 1].data_label
        dl.position = XL_LABEL_POSITION.ABOVE
        dl.text_frame.text = last_label
        style_run(dl.text_frame.paragraphs[0].runs[0], 8, True, TEAL)


def add_grouped_hbar(slide, spec, x, y, w, h):
    cd = CategoryChartData()
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["name"], s["values"])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    base_style(chart, legend=True)
    style_axes(chart, spec.get("axisFormat"))
    chart.category_axis.reverse_order = True
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.overlap = -10
    for s_spec, series in zip(spec["series"], plot.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(s_spec["color"])


def add_donut(slide, spec, x, y, w, h):
    cd = CategoryChartData()
    cd.categories = [truncate(c, 34) for c in spec["categories"]]
    cd.add_series("Nilai", spec["values"])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    base_style(chart, legend=True)
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    hole = chart.plots[0]._element.find(qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "58")
    plot = chart.plots[0]
    plot.has_data_labels = False
    for i, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(spec["colors"][i])
        point.format.line.color.rgb = WHITE


CHARTS = {"hbar": add_hbar, "line": add_line, "groupedHbar": add_grouped_hbar, "donut": add_donut}


def chart_panel(slide, spec, x, y, w, h):
    cx, cy, cw, ch = panel(slide, spec["title"], x, y, w, h)
    if not spec.get("categories"):
        text_box(slide, "Tidak ada data", cx, cy, cw, ch, 10, False, MUTED, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        return
    CHARTS[spec["kind"]](slide, spec, cx, cy, cw, ch)


# ----------------------------------------------------------------------------
# slides
# ----------------------------------------------------------------------------
def content_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[2])  # Title Only
    ph = get_ph(slide, 14)
    if ph is not None:
        ph.text = ""
        r = ph.text_frame.paragraphs[0].add_run()
        r.text = title
        style_run(r, 18, True, TEAL)
    return slide


def render_title(prs, s, theme):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_color, sub_color = TITLE_TEXT[theme]
    ph = get_ph(slide, 0)
    if ph is not None:
        ph.text = ""
        r = ph.text_frame.paragraphs[0].add_run()
        r.text = s["title"]
        style_run(r, 30, True, title_color)
    # Clear the template's subtitle placeholder and use our own textbox for
    # reliable position and contrast.
    for ph in slide.placeholders:
        if ph.placeholder_format.idx in (10, 11):
            ph.text = ""
    text_box(slide, s["subtitle"], 0.73, 4.05, 5.3, 0.8, 13, False, sub_color)  # stays left of the template graphic


def render_content(prs, s):
    slide = content_slide(prs, s["title"])
    y = TOP
    if s.get("banner"):
        y = banner(slide, s["banner"], y)
    y = kpi_strip(slide, s.get("kpis"), y)

    charts = s.get("charts") or []
    bullets = s.get("bullets") or []
    height = BOTTOM - y

    if len(charts) >= 3:
        # Three charts side by side, explanation as a strip underneath.
        explain_h = (0.5 + 0.21 * len(bullets)) if bullets else 0
        chart_h = height - (explain_h + GAP if bullets else 0)
        w = (RIGHT - LEFT - GAP * (len(charts) - 1)) / len(charts)
        for i, c in enumerate(charts):
            chart_panel(slide, c, LEFT + i * (w + GAP), y, w, chart_h)
        if bullets:
            explain_box(slide, "PENJELASAN", bullets, LEFT, y + chart_h + GAP, RIGHT - LEFT, explain_h, size=9, spacing=2)
        return

    explain_w = 3.45 if bullets else 0
    area_w = RIGHT - LEFT - (explain_w + GAP if bullets else 0)
    if charts:
        w = (area_w - GAP * (len(charts) - 1)) / len(charts)
        for i, c in enumerate(charts):
            chart_panel(slide, c, LEFT + i * (w + GAP), y, w, height)
    if bullets:
        explain_box(slide, "PENJELASAN", bullets, RIGHT - explain_w, y, explain_w, height)


def render_summary(prs, s):
    slide = content_slide(prs, s["title"])
    y = kpi_strip(slide, s.get("kpis"), TOP)
    height = BOTTOM - y
    paragraph = s.get("paragraph")
    bullets = s.get("bullets") or []
    if paragraph:
        w = (RIGHT - LEFT - GAP) / 2
        explain_box(slide, "ANALISIS", [], LEFT, y, w, height, paragraph=paragraph)
        explain_box(slide, s["bulletsTitle"], bullets, LEFT + w + GAP, y, w, height)
    else:
        # Size the panel to its content instead of leaving a mostly empty box.
        fit = min(height, 0.75 + 0.5 * max(1, len(bullets)))
        explain_box(slide, s["bulletsTitle"], bullets, LEFT, y, RIGHT - LEFT, fit, size=13 if len(bullets) <= 5 else 11)


def render(spec):
    theme = spec.get("theme") if spec.get("theme") in TEMPLATES else "black"
    prs = Presentation(TEMPLATES[theme])
    delete_all_slides(prs)
    for s in spec["slides"]:
        kind = s["type"]
        if kind == "title":
            render_title(prs, s, theme)
        elif kind == "content":
            render_content(prs, s)
        elif kind == "summary":
            render_summary(prs, s)
        elif kind == "end":
            prs.slides.add_slide(prs.slide_layouts[7])  # template carries "Terima Kasih"
    return prs


if __name__ == "__main__":
    spec = json.loads(sys.stdin.buffer.read().decode("utf-8"))
    out = sys.argv[1] if len(sys.argv) > 1 else "output.pptx"
    render(spec).save(out)
